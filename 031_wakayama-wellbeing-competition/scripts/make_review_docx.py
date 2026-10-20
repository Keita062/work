"""パーソルキャリア様_レビュー依頼_0904.md から送付用の Word 文書を生成する。

docx は手で編集せず、内容を変えたら md を直してこのスクリプトを再実行する
（make_gantt.py / make_review_slides.py と同じ運用）。

    python Wakayama/scripts/make_review_docx.py

入力: Wakayama/docs/パーソルキャリア様_レビュー依頼_0904.md
出力: Wakayama/パーソルキャリア様_レビュー依頼_0904.docx
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "パーソルキャリア様_レビュー依頼_0904.md"
OUT = ROOT / "パーソルキャリア様_レビュー依頼_0904.docx"

# ---------------------------------------------------------------- 体裁の定義

FONT = "Yu Gothic"
MONO = "MS Gothic"  # 全角も等幅になるのでフロー図用

NAVY = RGBColor(0x12, 0x34, 0x4D)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
ORANGE = RGBColor(0xC0, 0x56, 0x21)
INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x6B, 0x7A, 0x85)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

LINE_HEX = "D5DEE3"
PALE_HEX = "F2F6F8"
NAVY_HEX = "12344D"
WHITE_HEX = "FFFFFF"
QUOTE_HEX = "F7FAFB"

BODY_PT = 10.5
CONTENT_W = Cm(21.0 - 2.0 - 2.0)  # A4幅 − 左右マージン


# ---------------------------------------------------------------- 低レベル部品

def _font(run, size, *, bold=False, color=INK, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    # 日本語グリフにも同じフォントを当てる（latin だけだと游ゴシックにならない）
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


# OOXML は子要素の順序が決まっているので、正しい位置に差し込む
PPR_ORDER = [
    "pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
    "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
    "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
    "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
    "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
    "suppressOverlap", "jc", "textDirection", "textAlignment",
    "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
]
TBLPR_ORDER = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
    "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
]
TCPR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
]


def _insert_ordered(parent, el, order):
    name = el.tag.split("}")[1]
    idx = order.index(name)
    for child in parent:
        cname = child.tag.split("}")[1]
        if cname in order and order.index(cname) > idx:
            child.addprevious(el)
            return
    parent.append(el)


def _shade(cell_or_par, hex_color):
    """セル／段落に背景色を敷く。"""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    if hasattr(cell_or_par, "_tc"):
        _insert_ordered(cell_or_par._tc.get_or_add_tcPr(), shd, TCPR_ORDER)
    else:
        _insert_ordered(cell_or_par._p.get_or_add_pPr(), shd, PPR_ORDER)


def _borders(el, *, sides, color=LINE_HEX, sz=6, tag="w:tblBorders"):
    """表または段落に罫線を引く。sides は 'top left bottom right insideH insideV'。"""
    holder = OxmlElement(tag)
    for side in sides.split():
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color)
        holder.append(b)
    if tag == "w:tblBorders":
        _insert_ordered(el._tbl.tblPr, holder, TBLPR_ORDER)
    else:
        _insert_ordered(el._p.get_or_add_pPr(), holder, PPR_ORDER)


def _spacing(par, *, before=0, after=6, line=1.15):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def _hyperlink(par, text, url, size):
    """python-docx にネイティブAPIが無いので関係IDを自前で張る。"""
    rid = par.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = par.add_run()
    _font(run, size, color=TEAL)
    run.font.underline = True
    run.text = text
    link.append(run._r)
    par._p.append(link)


# ---------------------------------------------------------------- インライン記法

INLINE = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\([^)]+\))")


def add_runs(par, text, size, *, color=INK, bold=False, name=FONT):
    """**太字** と [表示文字](URL) を解釈しながら run を積む。"""
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            run = par.add_run(chunk[2:-2])
            _font(run, size, bold=True, color=color, name=name)
        elif chunk.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", chunk).groups()
            _hyperlink(par, label, url, size)
        else:
            for i, piece in enumerate(chunk.split("\n")):
                if i:
                    par.add_run().add_break()
                run = par.add_run(piece)
                _font(run, size, bold=bold, color=color, name=name)


# ---------------------------------------------------------------- ブロック描画

def h_title(doc, text, date_line):
    par = doc.add_paragraph()
    _spacing(par, after=4, line=1.2)
    run = par.add_run(text.replace(" — ", "\n"))
    _font(run, 20, bold=True, color=NAVY)
    par = doc.add_paragraph()
    _spacing(par, after=10)
    add_runs(par, date_line, 10, color=MUTED)
    rule(doc, color="0E7C86", sz=18)


def rule(doc, *, color=LINE_HEX, sz=6, after=10):
    par = doc.add_paragraph()
    _spacing(par, before=0, after=after, line=1.0)
    par.paragraph_format.space_before = Pt(2)
    _borders(par, sides="bottom", color=color, sz=sz, tag="w:pBdr")


def h1(doc, text):
    par = doc.add_paragraph()
    _spacing(par, before=14, after=6, line=1.2)
    run = par.add_run(text)
    _font(run, 15, bold=True, color=NAVY)
    _borders(par, sides="bottom", color=LINE_HEX, sz=8, tag="w:pBdr")


def h2(doc, text):
    par = doc.add_paragraph()
    _spacing(par, before=10, after=4, line=1.2)
    run = par.add_run(text)
    _font(run, 12, bold=True, color=TEAL)


def body(doc, text):
    par = doc.add_paragraph()
    _spacing(par, after=7, line=1.35)
    add_runs(par, text, BODY_PT)


def bullet(doc, text):
    par = doc.add_paragraph()
    _spacing(par, after=4, line=1.3)
    par.paragraph_format.left_indent = Cm(0.6)
    par.paragraph_format.first_line_indent = Cm(-0.35)
    run = par.add_run("・")
    _font(run, BODY_PT, color=TEAL, bold=True)
    add_runs(par, text, BODY_PT)


def callout(doc, text):
    """`> ## 〜` の一文を、枠つきの主張ブロックとして置く。"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = CONTENT_W
    _shade(cell, NAVY_HEX)
    _borders(tbl, sides="top left bottom right", color=NAVY_HEX, sz=8)
    par = cell.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(par, before=6, after=6, line=1.25)
    add_runs(par, text, 13.5, color=WHITE, bold=True)
    _gap(doc)


def quote(doc, lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = CONTENT_W
    _shade(cell, QUOTE_HEX)
    _borders(tbl, sides="top left bottom right", color=LINE_HEX, sz=6)
    first = True
    for text in lines:
        par = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        _spacing(par, before=2, after=2, line=1.3)
        if text:
            add_runs(par, text, 9.5, color=INK)
    _gap(doc)


def code(doc, lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = CONTENT_W
    _shade(cell, PALE_HEX)
    _borders(tbl, sides="top left bottom right", color=LINE_HEX, sz=6)
    first = True
    for text in lines:
        par = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        _spacing(par, before=1, after=1, line=1.2)
        run = par.add_run(text)
        _font(run, 9, color=NAVY, name=MONO)
    _gap(doc)


def _gap(doc, pt=8):
    par = doc.add_paragraph()
    _spacing(par, before=0, after=0, line=1.0)
    par.paragraph_format.space_after = Pt(pt)
    run = par.add_run("")
    _font(run, 1)


def _dw(text):
    """全角を2、半角を1として数えた表示幅。改行があれば最長行で測る。"""
    return max(
        sum(2 if ord(c) > 0x2000 else 1 for c in seg)
        for seg in (text.replace("<br>", "\n").split("\n") or [""])
    )


def _row_flags(row, *, header=False):
    """行を分割させない。ヘッダ行はページをまたぐとき繰り返す。"""
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))
    if header:
        tr_pr.append(OxmlElement("w:tblHeader"))


def table(doc, rows):
    """rows[0] をヘッダ行として描画する。"""
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]

    # 列幅は表示幅から按分する。長い列がすべてを飲み込まないよう指数で圧縮し、
    # 見出しは折り返させたくないので圧縮を弱くかける。
    weights = []
    for ci in range(ncol):
        content = max([_dw(r[ci]) for r in rows[1:]] or [0])
        head = _dw(rows[0][ci])
        weights.append(max(content ** 0.62, head ** 0.85, 2.0))
    total = sum(weights)

    tbl = doc.add_table(rows=len(rows), cols=ncol)
    tbl.autofit = False
    _borders(tbl, sides="top left bottom right insideH insideV",
             color=LINE_HEX, sz=6)

    for ri, row in enumerate(rows):
        _row_flags(tbl.rows[ri], header=(ri == 0))
        for ci, text in enumerate(row):
            cell = tbl.rows[ri].cells[ci]
            cell.width = int(CONTENT_W * weights[ci] / total)
            if ri == 0:
                _shade(cell, NAVY_HEX)
            else:
                _shade(cell, PALE_HEX if ri % 2 == 0 else WHITE_HEX)
            par = cell.paragraphs[0]
            _spacing(par, before=2, after=2, line=1.25)
            text = text.replace("<br>", "\n")
            if ri == 0:
                add_runs(par, text, 9.5, color=WHITE, bold=True)
            else:
                add_runs(par, text, 9.5)
    _gap(doc)


# ---------------------------------------------------------------- md の読み取り

def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_sep(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def build() -> Document:
    lines = SRC.read_text(encoding="utf-8").splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(1.8)
    sec.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    i = 0
    pending_title = None
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if not line.strip():
            i += 1
            continue

        # 見出し
        if line.startswith("# "):
            pending_title = line[2:].strip()
            i += 1
            continue
        if pending_title and line.startswith("**") and line.endswith("**"):
            h_title(doc, pending_title, line)
            pending_title = None
            i += 1
            continue
        if line.startswith("### "):
            h2(doc, line[4:].strip())
            i += 1
            continue
        if line.startswith("## "):
            h1(doc, line[3:].strip())
            i += 1
            continue

        # 水平線（見出し直前の区切りは h1 の下線に任せて捨てる）
        if line.strip() == "---":
            i += 1
            continue

        # フェンス付きコードブロック
        if line.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            code(doc, buf)
            i += 1
            continue

        # 引用
        if line.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i].lstrip(">").strip())
                i += 1
            if len(buf) == 1 and buf[0].startswith("## "):
                callout(doc, buf[0][3:].strip())
            else:
                quote(doc, buf)
            continue

        # 表
        if line.startswith("|"):
            buf = []
            while i < len(lines) and lines[i].startswith("|"):
                if not is_sep(lines[i]):
                    buf.append(split_row(lines[i]))
                i += 1
            table(doc, buf)
            continue

        # 箇条書き
        if line.startswith("- "):
            bullet(doc, line[2:].strip())
            i += 1
            continue

        body(doc, line.strip())
        i += 1

    footer(doc)
    return doc


def footer(doc):
    par = doc.sections[0].footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    _font(run, 9, color=MUTED)
    for instr in ("begin", "instrText", "end"):
        el = OxmlElement(f"w:fld{instr.capitalize()}" if instr != "instrText"
                         else "w:instrText")
        if instr == "instrText":
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        else:
            el.tag = qn("w:fldChar")
            el.set(qn("w:fldCharType"), instr)
        run._r.append(el)


def main():
    build().save(OUT)
    print(f"wrote {OUT.name}")


if __name__ == "__main__":
    main()
