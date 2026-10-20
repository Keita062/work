"""パーソルキャリア様_レビュー依頼_0904.md からレビュー会用スライドを生成する。

pptx は手で編集せず、内容を変えたらこのスクリプトを直して再生成する
（make_gantt.py と同じ運用）。

    python Wakayama/scripts/make_review_slides.py

出力: Wakayama/パーソルキャリア様_レビュー依頼_0904.pptx
"""

from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

# ---------------------------------------------------------------- 体裁の定義

OUT = Path(__file__).resolve().parents[1] / "パーソルキャリア様_レビュー依頼_0904.pptx"

FONT = "Yu Gothic"

NAVY = RGBColor(0x12, 0x34, 0x4D)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
ORANGE = RGBColor(0xC0, 0x56, 0x21)
INK = RGBColor(0x1F, 0x29, 0x33)
MUTED = RGBColor(0x6B, 0x7A, 0x85)
LINE = RGBColor(0xD5, 0xDE, 0xE3)
PALE = RGBColor(0xF2, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SW, SH = Emu(12192000), Emu(6858000)  # 16:9
ML = Emu(685800)  # 左マージン 0.75in
MR = Emu(685800)
CW = SW - ML - MR  # コンテンツ幅


def _font(run, size, *, bold=False, color=INK, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = FONT
    # 日本語グリフにも同じフォントを当てる（latin だけだと游ゴシックにならない）
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", FONT)


def _tb(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def _para(tf, text, size, *, bold=False, color=INK, space_after=6,
          align=PP_ALIGN.LEFT, first=False, line=1.25, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    run = p.add_run()
    run.text = text
    _font(run, size, bold=bold, color=color, italic=italic)
    return p


def _rect(slide, x, y, w, h, fill, *, line_color=None, shape=MSO_SHAPE.RECTANGLE):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(1)
    sp.text_frame.word_wrap = True
    return sp


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def slide(prs, kicker, title, *, sub=None):
    """見出し付きの標準スライド。本文の開始 y を返す。"""
    s = _blank(prs)
    _rect(s, Emu(0), Emu(0), SW, SH, WHITE)
    # 左肩のアクセントバー
    _rect(s, ML, Emu(430000), Emu(50000), Emu(430000), TEAL)

    tf = _tb(s, ML + Emu(150000), Emu(400000), CW, Emu(300000))
    _para(tf, kicker, 12, bold=True, color=TEAL, space_after=2, first=True)
    if title:
        _para(tf, title, 27, bold=True, color=NAVY, space_after=0)

    y = Emu(1000000) if title else Emu(800000)
    if sub:
        # 見出し（27pt）の下端を避けて置く。ここを詰めると本文と重なる
        tf2 = _tb(s, ML, y + Emu(160000), CW, Emu(300000))
        _para(tf2, sub, 14, color=MUTED, space_after=0, first=True)
        y += Emu(560000)
    else:
        y += Emu(160000)

    # 区切り線
    ln = _rect(s, ML, y - Emu(90000), CW, Emu(9000), LINE)
    ln.line.fill.background()
    return s, y + Emu(120000)


def table(slide_obj, x, y, w, cols, rows, *, header=True, size=13,
          head_size=12, row_h=Emu(430000), head_h=Emu(360000),
          zebra=True, bold_col0=False):
    """cols=列幅比のリスト, rows=行（セル文字列のリスト）のリスト。"""
    widths = [int(w * c / sum(cols)) for c in cols]
    cy = y
    for ri, row in enumerate(rows):
        is_head = header and ri == 0
        h = head_h if is_head else row_h
        cx = x
        for ci, cell in enumerate(row):
            if is_head:
                fill = NAVY
            elif zebra and (ri % 2 == 0):
                fill = PALE
            else:
                fill = WHITE
            box = _rect(slide_obj, cx, cy, widths[ci], h, fill, line_color=LINE)
            tf = box.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Emu(90000)
            tf.margin_top = tf.margin_bottom = Emu(36000)
            _para(
                tf, cell,
                head_size if is_head else size,
                bold=is_head or (bold_col0 and ci == 0),
                color=WHITE if is_head else INK,
                space_after=0, first=True, line=1.15,
                align=PP_ALIGN.LEFT,
            )
            cx += widths[ci]
        cy += h
    return cy


def footer(prs):
    """表紙以外に通しノンブルを振る。"""
    for i, s in enumerate(prs.slides):
        if i == 0:
            continue
        tf = _tb(s, SW - MR - Emu(700000), SH - Emu(560000), Emu(700000), Emu(280000))
        _para(tf, str(i), 10, color=MUTED, space_after=0, first=True,
              align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------- 各スライド

def build() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH

    # ---------------------------------------------------------- 1 表紙
    s = _blank(prs)
    _rect(s, Emu(0), Emu(0), SW, SH, NAVY)
    _rect(s, Emu(0), SH - Emu(180000), SW, Emu(180000), TEAL)
    tf = _tb(s, ML, Emu(2100000), CW, Emu(2000000))
    _para(tf, "第10回 和歌山県データ利活用コンペティション（大学生部門）",
          15, color=RGBColor(0x9F, 0xC4, 0xD0), space_after=16, first=True)
    _para(tf, "和歌山県 主観的ウェルビーイング分析", 40, bold=True,
          color=WHITE, space_after=8)
    _para(tf, "分析設計レビュー会", 26, color=RGBColor(0x7F, 0xD4, 0xDE),
          space_after=34)
    _para(tf, "2026年9月4日　鈴木啓太・西村瞭一", 15,
          color=RGBColor(0xC9, 0xD8, 0xDF), space_after=0)

    # ---------------------------------------------------------- 2 サマリー
    s, y = slide(prs, "EXECUTIVE SUMMARY", "本日お時間をいただいた理由")
    box = _rect(s, ML, y, CW, Emu(1520000), PALE, line_color=None)
    bar = _rect(s, ML, y, Emu(50000), Emu(1520000), TEAL)
    bar.line.fill.background()
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(280000)
    tf.margin_right = Emu(220000)
    _para(tf, "和歌山県の住民の幸福度を上げるために、どの分野に、どんな手を打つべきかを"
              "明らかにする設計を組み、いま実装に入る直前です。",
          17, bold=True, color=NAVY, space_after=10, first=True)
    _para(tf, "この設計を一通りご覧いただき、全体の筋から各ステップの狙いまで、"
              "気になった点をご指摘いただきたく、お時間をいただきました。",
          17, color=INK, space_after=0)

    y += Emu(1780000)
    table(s, ML, y, CW, [1.6, 4.4, 4.0], [
        ["", "いまの状態", "本日いただきたいもの"],
        ["設計", "2名で議論を重ね、一旦完成", "全体の筋がずれていないかのご指摘"],
        ["実装", "これから着手（コード・データはまだ0本）", "「実務ならここで止まる」という箇所"],
        ["締切", "応募 2026-10-08 ／ 最終審査会 12-12", "優先順位についてのご助言"],
    ], size=14, row_h=Emu(500000))

    # ---------------------------------------------------------- 3 アジェンダ
    s, y = slide(prs, "AGENDA", "本日の進め方")
    items = [
        ("1", "本レビュー会の目的"),
        ("2", "本会議の背景"),
        ("3", "コンペの議題と募集テーマ"),
        ("4", "分析の目的"),
        ("5", "目的を叶えるための方法"),
        ("6", "収集するデータ"),
        ("7", "分析設計"),
    ]
    col_w = int(CW / 2) - Emu(200000)
    for i, (num, label) in enumerate(items):
        cx = ML + (0 if i < 4 else int(CW / 2) + Emu(200000))
        cy = y + Emu(180000) + (i % 4) * Emu(920000)
        chip = _rect(s, cx, cy, Emu(560000), Emu(560000), NAVY)
        ctf = chip.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _para(ctf, num, 20, bold=True, color=WHITE, space_after=0,
              first=True, align=PP_ALIGN.CENTER)
        tf = _tb(s, cx + Emu(760000), cy + Emu(130000),
                 col_w - Emu(760000), Emu(400000))
        _para(tf, label, 19, bold=True, color=INK, space_after=0, first=True)

    # ---------------------------------------------------------- 4 §1 目的
    s, y = slide(prs, "1 ／ 本レビュー会の目的", "実装前のいま、外部の目を入れたい")
    box = _rect(s, ML, y + Emu(300000), CW, Emu(1150000), NAVY)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, "実装に入る前に、この設計で目的に到達できるかを確認すること。",
          26, bold=True, color=WHITE, space_after=0, first=True,
          align=PP_ALIGN.CENTER)

    tf = _tb(s, ML, y + Emu(1750000), CW, Emu(1400000))
    _para(tf, "実装前のこの段階であれば、いただいたご指摘を設計に反映できます。",
          18, color=INK, space_after=16, first=True)
    _para(tf, "全体の筋から各ステップの狙いまで、幅広くご指摘いただけますと幸いです。",
          18, color=INK, space_after=24)
    _para(tf, "説明の途中で止めていただいても構いません。"
              "「この前提が違う」「実務ではこうする」という一言が一番ありがたいです。",
          16, color=ORANGE, bold=True, space_after=0)

    # ---------------------------------------------------------- 5 §2 背景
    s, y = slide(prs, "2 ／ 本会議の背景", "分析設計を一旦完成させたところです")
    tf = _tb(s, ML, y, CW, Emu(500000))
    _para(tf, "2名で議論を重ね、設計を組み終えました。実装はこれからで、"
              "本日のレビューを反映したうえで着手します。",
          17, color=INK, space_after=0, first=True)

    y += Emu(750000)
    table(s, ML, y, int(CW * 0.72), [3.2, 2.0], [
        ["今後の予定", "日程"],
        ["応募締切", "2026-10-08（木）"],
        ["最終審査会（プレゼン）", "2026-12-12（土）"],
    ], size=16, row_h=Emu(520000), head_h=Emu(430000), bold_col0=True)

    # ---------------------------------------------------------- 6 §3-1 概要
    s, y = slide(prs, "3-1 ／ コンペの議題", "第10回 和歌山県データ利活用コンペティション")
    table(s, ML, y, CW, [2.3, 7.7], [
        ["項目", "内容"],
        ["募集テーマ", "「データで描くウェルビーング　地域の幸福度を高めるためのまちづくり」"],
        ["主催 ／ 部門", "和歌山県 ／ 大学生部門"],
        ["応募締切", "2026-10-08（木）"],
        ["提出物", "応募用紙（Word/PDF）＋ プレゼン資料（PowerPoint/PDF・10MB未満）"],
        ["審査の流れ", "事前審査（形式）→ 一次審査（書類・11月末までに通知）→ 最終審査（プレゼン）"],
        ["最終審査会", "2026-12-12（土）10:00〜／和歌山県データ利活用推進センター"],
        ["プレゼン時間", "大学生部門 14分"],
        ["賞", "大賞20万円／データ利活用賞5万円／政策アイデア賞5万円（商品券・各部門1点）"],
    ], size=13, row_h=Emu(420000), head_h=Emu(360000), bold_col0=True)

    # ---------------------------------------------------------- 7 §3-2 審査観点
    s, y = slide(prs, "3-2 ／ 審査の観点", "この2点で審査されます（募集要項より）")
    half = int((CW - Emu(300000)) / 2)
    for i, (head, body) in enumerate([
        ("ア：データ利活用の観点",
         "データを適切に利活用し、独自の着眼点からデータを分析し、"
         "地域の現状や課題を抽出できているか"),
        ("イ：施策提案の観点",
         "募集テーマに基づいた独自性・将来性のある施策提案であり、"
         "データ分析から得られたエビデンス（客観的証拠）に基づいているか"),
    ]):
        cx = ML + i * (half + Emu(300000))
        card = _rect(s, cx, y, half, Emu(1750000), PALE, line_color=LINE)
        top = _rect(s, cx, y, half, Emu(80000), TEAL)
        top.line.fill.background()
        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Emu(280000)
        _para(tf, head, 19, bold=True, color=NAVY, space_after=14, first=True)
        _para(tf, body, 15, color=INK, space_after=0)

    y += Emu(2050000)
    box = _rect(s, ML, y, CW, Emu(620000), RGBColor(0xFD, 0xF1, 0xE7),
                line_color=ORANGE)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(280000)
    _para(tf, "加えて「応募書類の不備や誤りなど、資料の完成度も審査します」"
              "と要項に明記されています。",
          16, bold=True, color=ORANGE, space_after=0, first=True)

    # ---------------------------------------------------------- 8 §4 目的
    s, y = slide(prs, "4 ／ 分析の目的", None)
    box = _rect(s, ML, y + Emu(500000), CW, Emu(1600000), NAVY)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(400000)
    _para(tf, "和歌山県の住民が実際に感じている幸福度\n（主観的ウェルビーイング）を上げること。",
          30, bold=True, color=WHITE, space_after=0, first=True,
          align=PP_ALIGN.CENTER, line=1.4)

    tf = _tb(s, ML, y + Emu(2350000), CW, Emu(400000))
    _para(tf, "※ 目的変数は「生活満足度」ではなく「幸福度」に置きます。"
              "住民が実際に感じている幸福を上げることが目的であり、"
              "募集テーマの言葉とも一致するためです。",
          14, color=MUTED, space_after=0, first=True)

    # ---------------------------------------------------------- 9 §5 方法
    s, y = slide(prs, "5 ／ 目的を叶えるための方法", None)
    box = _rect(s, ML, y, CW, Emu(950000), TEAL)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, "分析から政策提言までを行う。", 30, bold=True, color=WHITE,
          space_after=0, first=True, align=PP_ALIGN.CENTER)

    y += Emu(1200000)
    tf = _tb(s, ML, y, CW, Emu(1400000))
    _para(tf, "理由", 15, bold=True, color=TEAL, space_after=10, first=True)
    _para(tf, "目的を叶えるには、データドリブンで幸福度を向上させる施策が"
              "何なのかを明確にし、手を打つことが必要だからです。",
          18, color=INK, space_after=10)
    _para(tf, "施策が明確になっていなければ手は打てず、手を打たなければ幸福度は上がりません。"
              "したがって、分析だけでも提言だけでもなく、"
              "分析から政策提言までを一続きで行うことを方法とします。",
          18, color=INK, space_after=0)

    # ---------------------------------------------------------- 10 §5 到達状態
    s, y = slide(prs, "5 ／ この方法をとった後の状態",
                 "分析から政策提言までを終えたとき、次の4点を言える状態を目指します")
    table(s, ML, y, CW, [0.5, 5.2, 4.3], [
        ["#", "言えるようにすること", "それが言えると何ができるか"],
        ["1", "どの分野に手を打つべきか（優先順位と、その根拠）", "予算を振り向ける先が決まる"],
        ["2", "その分野は、周知・運用で実感を上げられるのか／\n環境そのものを変えないと上がらないのか",
         "打つべき施策の型（ソフト／ハード）が決まる"],
        ["3", "環境を変えるなら、どの指標を、どこまで動かすのか", "施策のターゲットと目標値が決まる"],
        ["4", "その打ち手で、幸福度がどれだけ動くと見込めるのか", "費用対効果を示して意思決定にかけられる"],
    ], size=14, row_h=Emu(600000))

    tf = _tb(s, ML, Emu(5850000), CW, Emu(400000))
    _para(tf, "この4点が埋まれば「和歌山県はここに、こう手を打つべき」という形で提言に落ちます。"
              "逆に、ここが埋まらない設計は目的を達成しません。",
          15, bold=True, color=ORANGE, space_after=0, first=True)

    # ---------------------------------------------------------- 11 §6-1/2 データ
    s, y = slide(prs, "6-1 ／ 使うデータ", "デジタル庁「地域幸福度（Well-Being）指標」ダッシュボード")
    box = _rect(s, ML, y, CW, Emu(760000), PALE, line_color=LINE)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(280000)
    tf.margin_right = Emu(280000)
    _para(tf, "市区町村単位で、住民の実感（主観指標）と地域の環境水準（客観指標）が、"
              "同じ24分野の枠組みで並んで取れる。全国約530自治体分が公開。",
          16, bold=True, color=NAVY, space_after=0, first=True)

    y += Emu(1000000)
    table(s, ML, y, CW, [3.4, 6.6], [
        ["取れるもの", "分かること"],
        ["幸福度（0〜10の総合スコア）", "住民が全体としてどれだけ幸せを感じているか　→ これを目的変数にする"],
        ["主観24分野（住民の実感）", "24分野それぞれについて、住民が実際にどう感じているか"],
        ["客観24分野（環境の水準）", "同じ24分野について、地域の環境が客観的にどの水準にあるか"],
        ["客観の個別KPI（208本）", "客観24分野が、どの具体的な指標から成り立っているか"],
        ["年代別の内訳", "主観側は全分野が年代別に取れる　→ 年代によって実感がどう違うかを見られる"],
    ], size=14, row_h=Emu(490000), bold_col0=True)

    tf = _tb(s, ML, Emu(5750000), CW, Emu(500000))
    _para(tf, "主観と客観が同じ枠組みで並ぶため、"
              "「環境はあるのに、住民の実感が伴っていない分野」を特定できる。"
              "これが本分析の起点です。",
          16, bold=True, color=ORANGE, space_after=0, first=True)

    # ---------------------------------------------------------- 12 24分野
    s, y = slide(prs, "6-2 ／ 24分野の内訳", None)
    groups = [
        ("生活環境（16）", "医療・福祉／買物・飲食／住宅環境／移動・交通／遊び・娯楽／子育て／"
                        "初等・中等教育／地域行政／デジタル生活／公共空間／都市景観／自然景観／"
                        "自然の恵み／環境共生／自然災害／事故・犯罪"),
        ("地域の人間関係（2）", "地域とのつながり／多様性と寛容性"),
        ("自分らしい生き方（6）", "自己効力感／健康状態／文化・芸術／教育機会の豊かさ／雇用・所得／事業創造"),
    ]
    heights = [Emu(1300000), Emu(800000), Emu(950000)]
    cy = y + Emu(150000)
    for (head, body), h in zip(groups, heights):
        bar = _rect(s, ML, cy, Emu(50000), h, TEAL)
        bar.line.fill.background()
        card = _rect(s, ML + Emu(50000), cy, CW - Emu(50000), h, PALE,
                     line_color=None)
        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Emu(280000)
        tf.margin_right = Emu(280000)
        _para(tf, head, 17, bold=True, color=NAVY, space_after=10, first=True)
        _para(tf, body, 15, color=INK, space_after=0)
        cy += h + Emu(180000)

    # ---------------------------------------------------------- 13 §6-3 出所
    s, y = slide(prs, "6-3 ／ データの出所", "主指標はすべて住民アンケート、客観指標は官庁統計からの合成")
    table(s, ML, y, CW, [2.6, 7.4], [
        ["区分", "出所"],
        ["主指標（主観）", "すべて住民アンケート。デジタル庁が全国の自治体で実施"],
        ["客観指標",
         "官庁統計から合成。国交省「都市モニタリングシート」、厚労省「医療費の地域差分析」、\n"
         "総務省「社会・人口統計体系」等。原出典まで公開されている"],
        ["環境変数（統制用）",
         "総務省（人口・高齢化率・財政力指数）。\n"
         "ダッシュボードの外から足す変数で、自治体間の条件差を吸収するために置く"],
    ], size=15, row_h=Emu(640000), bold_col0=True)

    # ---------------------------------------------------------- 14 §6-4 分析単位
    s, y = slide(prs, "6-4 ／ 目的変数と分析単位", None)
    table(s, ML, y, CW, [2.4, 7.6], [
        ["項目", "内容"],
        ["目的変数", "幸福度（0〜10の生値）。コンペのテーマ名に一致させる"],
        ["説明変数",
         "主観24分野・客観24分野（いずれも偏差値）、客観の個別KPI（実数）、\n"
         "環境変数（人口・高齢化率・財政力指数などの地域条件。統制変数として置く）"],
        ["分析単位", "1行 ＝ 1自治体 × 1年度 × 1年代"],
        ["対象",
         "推定は全国の自治体、診断の対象は和歌山。\n"
         "全国データで「同水準の自治体ならこのくらい」という物差しを作り、それを和歌山に当てる"],
    ], size=14, row_h=Emu(680000), bold_col0=True)

    # ---------------------------------------------------------- 15 §7-1 用語
    s, y = slide(prs, "7-1 ／ この資料で使う言葉",
                 "分析設計の説明に入る前に、以下の5語を定義します")
    table(s, ML, y, CW, [2.6, 7.4], [
        ["言葉", "指すもの"],
        ["幸福度", "目的変数。住民に直接たずねた0〜10の総合スコア"],
        ["主観スコア", "24分野それぞれの、住民の実感の値（偏差値）"],
        ["客観スコア", "24分野それぞれの、地域の環境水準の値（偏差値）"],
        ["客観的構成値", "客観スコアを組み立てている個別KPI（24分野あわせて208本）"],
        ["追加要因",
         "客観的構成値に含まれていないが、主観スコアに関係している可能性のある要因\n"
         "（情報の届き方、予約のしやすさ、運用など）"],
    ], size=14, row_h=Emu(570000), bold_col0=True)

    # ---------------------------------------------------------- 16 §7-2 流れ
    s, y = slide(prs, "7-2 ／ 全体の流れ", "分析は9ステップですが、かたまりで見ると4つです")
    blocks = [
        ("①", "どの分野を扱うかを決める", "STEP 1・2", NAVY),
        ("②", "その分野の実感が、環境の数字で説明できるかを判定する", "STEP 3・4", TEAL),
        ("③", "何をどこまで動かすと、実感がどれだけ動くかを出す", "STEP 5・6・7", NAVY),
        ("④", "行政が打てる手に翻訳する", "STEP 8・9", TEAL),
    ]
    bh = Emu(810000)
    gap = Emu(300000)
    cy = y + Emu(120000)
    for i, (num, label, steps, color) in enumerate(blocks):
        chip = _rect(s, ML, cy, Emu(830000), bh, color)
        ctf = chip.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _para(ctf, num, 26, bold=True, color=WHITE, space_after=0, first=True,
              align=PP_ALIGN.CENTER)
        body = _rect(s, ML + Emu(830000), cy, CW - Emu(830000) - Emu(1900000),
                     bh, PALE, line_color=None)
        btf = body.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        btf.margin_left = Emu(320000)
        _para(btf, label, 19, bold=True, color=INK, space_after=0, first=True)
        tag = _rect(s, ML + CW - Emu(1900000), cy, Emu(1900000), bh, WHITE,
                    line_color=LINE)
        ttf = tag.text_frame
        ttf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _para(ttf, steps, 16, bold=True, color=color, space_after=0, first=True,
              align=PP_ALIGN.CENTER)
        if i < 3:
            ar = _rect(s, ML + Emu(300000), cy + bh + Emu(20000),
                       Emu(230000), gap - Emu(40000), LINE,
                       shape=MSO_SHAPE.DOWN_ARROW)
            ar.line.fill.background()
        if i == 1:
            note = _tb(s, ML + Emu(700000), cy + bh + Emu(40000),
                       Emu(4600000), Emu(240000))
            _para(note, "ここで道が2つに分かれる", 12, bold=True, color=ORANGE,
                  space_after=0, first=True)
        cy += bh + gap

    # ---------------------------------------------------------- 17 9ステップ前半
    s, y = slide(prs, "7-2 ／ 9ステップ（①分野を絞る・②道を分ける）", None)
    table(s, ML, y, CW, [1.3, 0.7, 4.6, 3.4], [
        ["", "STEP", "やること", "ここで分かること"],
        ["① 分野を絞る", "1", "24分野の主観スコアと幸福度の関連性を確認する",
         "幸福度と関連が大きい分野はどれか"],
        ["", "2", "重点カテゴリーを選定する", "そのうち、和歌山が優先的に扱うべき分野はどれか"],
        ["② 道を分ける", "3", "重点カテゴリーの主観スコアを、客観的構成値で説明できるかを確認する",
         "環境の数字で説明がつく分野か、つかない分野か"],
        ["", "4A", "【説明できる場合】客観的構成値をそのまま改善候補として扱う",
         "改善候補（既存のKPI）"],
        ["", "4B", "【説明できない場合】追加要因を探索し、改善候補に加える",
         "改善候補（既存のKPI ＋ 追加要因）"],
    ], size=13, row_h=Emu(640000), bold_col0=True)

    # ---------------------------------------------------------- 18 9ステップ後半
    s, y = slide(prs, "7-2 ／ 9ステップ（③効き目を見積もる・④施策にする）", None)
    table(s, ML, y, CW, [1.3, 0.7, 4.6, 3.4], [
        ["", "STEP", "やること", "ここで分かること"],
        ["③ 効き目を\n　 見積もる", "5", "改善候補のうち、伸ばすと実感が動きそうなものに絞り込む",
         "効きそうな候補はどれか"],
        ["", "6", "和歌山に、その数値を伸ばす余地があるかを確認する", "実際に動かせる候補はどれか"],
        ["", "7", "候補に目標値を置き、主観スコアの変化を試算する",
         "どこまで動かすと、実感がどれだけ動くか"],
        ["④ 施策にする", "8", "行政施策の介入可能性と時間軸を整理する",
         "短期・中期・長期のどれで打てる手か"],
        ["", "9", "政策提案とロードマップを作成する", "和歌山が、何を、いつやるべきか"],
    ], size=13, row_h=Emu(620000), bold_col0=True)

    tf = _tb(s, ML, Emu(5950000), CW, Emu(300000))
    _para(tf, "※ 改善候補 ＝ STEP 4A では「客観的構成値」、"
              "STEP 4B では「客観的構成値」＋探索した「追加要因」",
          13, color=MUTED, space_after=0, first=True)

    # ---------------------------------------------------------- 19 §7-3 分岐
    s, y = slide(prs, "7-3 ／ 設計の要は STEP 3 の分岐",
                 "同じ「主観スコアが低い」でも、背景は2通りあります")
    y = table(s, ML, y, CW, [1.4, 4.3, 4.3], [
        ["", "状況", "意味すること"],
        ["ケースA", "客観スコアが低く、主観スコアも低い", "環境そのものが足りていない可能性がある"],
        ["ケースB", "客観スコアは一定水準だが、主観スコアは低い",
         "既存の客観指標では、実感の低さを説明しきれていない可能性がある"],
    ], size=14, row_h=Emu(590000), bold_col0=True)

    y += Emu(200000)
    box = _rect(s, ML, y, CW, Emu(500000), RGBColor(0xFD, 0xF1, 0xE7),
                line_color=ORANGE)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(260000)
    _para(tf, "この2つを分けずに進めると、打ち手が「施設を増やす」一辺倒になり、"
              "和歌山が実行できない提言になります。",
          15, bold=True, color=ORANGE, space_after=0, first=True)

    y += Emu(700000)
    table(s, ML, y, CW, [1.4, 1.8, 3.6, 3.2], [
        ["ルート", "判定", "次にやること", "想定される政策の型"],
        ["ルートA", "説明できる", "STEP 4A：個別KPI（客観的構成値）まで掘り下げる",
         "ハード寄り（環境の水準そのものを上げる）"],
        ["ルートB", "説明できない", "STEP 4B：客観環境が同程度の自治体を集めて比較し、追加要因を探す",
         "ソフト寄り（周知・運用・使いやすさを上げる）"],
    ], size=13, row_h=Emu(560000), bold_col0=True)

    # ---------------------------------------------------------- 20 クロージング
    s = _blank(prs)
    _rect(s, Emu(0), Emu(0), SW, SH, NAVY)
    _rect(s, Emu(0), SH - Emu(180000), SW, Emu(180000), TEAL)
    tf = _tb(s, ML, Emu(2700000), CW, Emu(1400000))
    _para(tf, "本日はお時間をいただき、ありがとうございます。", 30, bold=True,
          color=WHITE, space_after=20, first=True, align=PP_ALIGN.CENTER)
    _para(tf, "いただいたご指摘は、実装前の設計に反映いたします。", 18,
          color=RGBColor(0xC9, 0xD8, 0xDF), space_after=0, align=PP_ALIGN.CENTER)

    footer(prs)
    return prs


if __name__ == "__main__":
    prs = build()
    prs.save(OUT)
    print(f"wrote {OUT.name}  ({len(prs.slides._sldIdLst)} slides)")
