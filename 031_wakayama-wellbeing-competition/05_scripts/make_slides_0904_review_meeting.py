"""西村氏版のレビュー会資料 md からレビュー会用スライドを生成する。

出典: 共同リポジトリ wakayama_competition / origin/analysis_nishimura
      Nishimura/docs/0904和歌山県 主観的ウェルビーイング分析 — レビュー会資料.md
本リポジトリでの正本: 031_wakayama-wellbeing-competition/01_docs/0904和歌山県 主観的ウェルビーイング分析 — レビュー会資料.md

pptx は手で編集せず、内容を変えたらこのスクリプトを直して再生成する
（make_review_slides.py / make_gantt.py と同じ運用）。

    python 031_wakayama-wellbeing-competition/05_scripts/make_slides_0904_review_meeting.py

出力: 031_wakayama-wellbeing-competition/01_docs/0904_和歌山県主観的ウェルビーイング分析_レビュー会資料.pptx

体裁（色・フォント・表・見出し）は make_review_slides.py を共用する。
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

import make_review_slides as base  # noqa: E402
from make_review_slides import (  # 体裁の定義とレイアウト部品を共用
    CW,
    LINE,
    ML,
    NAVY,
    ORANGE,
    PALE,
    SH,
    SW,
    TEAL,
    WHITE,
    _blank,
    _para,
    _rect,
    _tb,
    footer,
    table,
)

OUT = (
    Path(__file__).resolve().parents[1]
    / "01_docs"
    / "0904_和歌山県主観的ウェルビーイング分析_レビュー会資料.pptx"
)

CREAM = RGBColor(0xFD, 0xF1, 0xE7)  # 注意書きボックスの地色

# --- 体裁の上書き ---------------------------------------------------------
# 游ゴシックは投影すると細くて薄く見えるため、線の太いメイリオにする。
# （「游ゴシック Medium」は Bold を持たないので、太字が合成されて二重に
#   描画される。太字を正式に持つ書体でないと見出しが汚くなる）
# 本文色・補足色も一段濃くする。make_review_slides 側の定数も差し替えて、
# 共用している table() / footer() にも効かせる。
FONT = "Meiryo"
INK = RGBColor(0x0F, 0x17, 0x1F)
MUTED = RGBColor(0x45, 0x53, 0x5D)
base.FONT, base.INK, base.MUTED = FONT, INK, MUTED


def slide(prs, kicker, title, *, sub=None):
    """見出し付きの標準スライド。本文の開始 y を返す。

    make_review_slides.slide() と同じ体裁だが、見出しを上に、区切り線を下に
    ずらして間隔を広げてある（元の配置だと見出しが線に接して見える）。
    """
    s = _blank(prs)
    _rect(s, Emu(0), Emu(0), SW, SH, WHITE)
    # 左肩のアクセントバー
    _rect(s, ML, Emu(310000), Emu(50000), Emu(430000), TEAL)

    tf = _tb(s, ML + Emu(150000), Emu(280000), CW, Emu(300000))
    _para(tf, kicker, 12, bold=True, color=TEAL, space_after=2, first=True)
    if title:
        _para(tf, title, 27, bold=True, color=NAVY, space_after=0)

    y = Emu(1080000) if title else Emu(880000)
    if sub:
        # 見出し（27pt）の下端を避けて置く。ここを詰めると本文と重なる
        tf2 = _tb(s, ML, y + Emu(160000), CW, Emu(300000))
        _para(tf2, sub, 14, color=MUTED, space_after=0, first=True)
        y += Emu(600000)
    else:
        y += Emu(240000)

    # 区切り線
    ln = _rect(s, ML, y - Emu(90000), CW, Emu(9000), LINE)
    ln.line.fill.background()
    return s, y + Emu(120000)


def bullets(s, y, items, *, size=19, gap=Emu(560000)):
    """行頭にティールの小さな四角を置いた箇条書き。次の y を返す。"""
    for i, text in enumerate(items):
        cy = y + i * gap
        mk = _rect(s, ML + Emu(50000), cy + Emu(95000),
                   Emu(140000), Emu(140000), TEAL)
        mk.line.fill.background()
        tf = _tb(s, ML + Emu(340000), cy, CW - Emu(340000), gap)
        _para(tf, text, size, color=INK, space_after=0, first=True)
    return y + len(items) * gap


def note_box(s, y, text, *, h=Emu(560000), size=15):
    """オレンジ枠の注意書き。次の y を返す。"""
    box = _rect(s, ML, y, CW, h, CREAM, line_color=ORANGE)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(260000)
    _para(tf, text, size, bold=True, color=ORANGE, space_after=0, first=True)
    return y + h


# ---------------------------------------------------------------- 各スライド

def build() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH

    # ---------------------------------------------------------- 1 表紙
    s = _blank(prs)
    _rect(s, Emu(0), Emu(0), SW, SH, NAVY)
    _rect(s, Emu(0), SH - Emu(180000), SW, Emu(180000), TEAL)
    tf = _tb(s, ML, Emu(2100000), CW, Emu(2000000))
    _para(tf, "第10回 和歌山県データ利活用コンペティション（大学生部門）",
          15, color=RGBColor(0x9F, 0xC4, 0xD0), space_after=16, first=True)
    _para(tf, "和歌山県 主観的ウェルビーイング分析", 40, bold=True,
          color=WHITE, space_after=8)
    _para(tf, "レビュー会資料", 26, color=RGBColor(0x7F, 0xD4, 0xDE),
          space_after=34)
    _para(tf, "2026年9月4日　鈴木啓太・西村瞭一", 15,
          color=RGBColor(0xC9, 0xD8, 0xDF), space_after=0)

    # ---------------------------------------------------------- 2 サマリー
    s, y = slide(prs, "EXECUTIVE SUMMARY", "本日お時間をいただいた理由")
    box = _rect(s, ML, y, CW, Emu(1700000), PALE, line_color=None)
    bar = _rect(s, ML, y, Emu(50000), Emu(1700000), TEAL)
    bar.line.fill.background()
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Emu(280000)
    tf.margin_right = Emu(220000)
    _para(tf, "住民の幸福度を上げるために、どの分野に、どんな手を打つべきか。\n"
              "それを明らかにする設計を組み終え、実装に入る直前です。",
          19, bold=True, color=NAVY, space_after=12, first=True)
    _para(tf, "本日は設計を一通りご覧いただき、気になった点をご指摘ください。",
          17, color=INK, space_after=0)

    y += Emu(1960000)
    table(s, ML, y, CW, [1.6, 4.4, 4.0], [
        ["", "いまの状態", "本日いただきたいもの"],
        ["設計", "2名で議論を重ね、一旦完成", "全体の筋がずれていないかのご指摘"],
        ["実装", "これから着手", "「実務ならここで止まる」という箇所"],
        ["締切", "応募 2026-10-08 ／ 最終審査会 12-12", "優先順位についてのご助言"],
    ], size=14, row_h=Emu(480000))

    # ---------------------------------------------------------- 3 アジェンダ
    s, y = slide(prs, "AGENDA", "本日の進め方")
    items = [
        ("1", "本レビュー会の目的"),
        ("2", "本会議の背景"),
        ("3", "コンペの議題と募集テーマ"),
        ("4", "分析の目的"),
        ("5", "目的を叶えるための方法"),
        ("6", "収集するデータ"),
        ("7", "分析設計"),
    ]
    col_w = int(CW / 2) - Emu(200000)
    for i, (num, label) in enumerate(items):
        cx = ML + (0 if i < 4 else int(CW / 2) + Emu(200000))
        cy = y + Emu(180000) + (i % 4) * Emu(920000)
        chip = _rect(s, cx, cy, Emu(560000), Emu(560000), NAVY)
        ctf = chip.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _para(ctf, num, 20, bold=True, color=WHITE, space_after=0,
              first=True, align=PP_ALIGN.CENTER)
        tf = _tb(s, cx + Emu(760000), cy + Emu(130000),
                 col_w - Emu(760000), Emu(400000))
        _para(tf, label, 19, bold=True, color=INK, space_after=0, first=True)

    # ---------------------------------------------------------- 4 §1 目的
    s, y = slide(prs, "1 ／ 本レビュー会の目的", "実装前のいま、外部の目を入れたい")
    box = _rect(s, ML, y + Emu(300000), CW, Emu(1150000), NAVY)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, "実装に入る前に、この設計で目的に到達できるかを\n外部レビューで確認すること。",
          25, bold=True, color=WHITE, space_after=0, first=True,
          align=PP_ALIGN.CENTER, line=1.35)

    tf = _tb(s, ML, y + Emu(1800000), CW, Emu(1400000))
    _para(tf, "いまなら、いただいたご指摘を設計に反映できます。",
          19, color=INK, space_after=22, first=True)
    _para(tf, "「この前提が違う」「実務ではこうする」の一言が一番ありがたいです。",
          17, color=ORANGE, bold=True, space_after=0)

    # ---------------------------------------------------------- 5 §2 背景
    s, y = slide(prs, "2 ／ 本会議の背景", "分析設計を一旦完成させたところです")
    tf = _tb(s, ML, y, CW, Emu(500000))
    _para(tf, "2名で議論を重ね、分析設計を一旦完成させました。実装はこれからで、"
              "本日のレビューを反映したうえで着手します。",
          17, color=INK, space_after=0, first=True)

    y += Emu(750000)
    table(s, ML, y, int(CW * 0.72), [3.2, 2.0], [
        ["今後の予定", "日程"],
        ["応募締切", "2026-10-08（木）"],
        ["最終審査会（プレゼン）", "2026-12-12（土）"],
    ], size=16, row_h=Emu(520000), head_h=Emu(430000), bold_col0=True)

    # ---------------------------------------------------------- 6 §3-1 概要
    s, y = slide(prs, "3-1 ／ コンペの議題と募集テーマ",
                 "第10回 和歌山県データ利活用コンペティション")
    table(s, ML, y, CW, [2.3, 7.7], [
        ["項目", "内容"],
        ["主催 ／ 部門", "和歌山県 ／ 大学生部門"],
        ["募集テーマ", "データで描くウェルビーング　地域の幸福度を高めるためのまちづくり"],
        ["応募締切", "2026-10-08（木）"],
        ["提出物", "応募用紙（Word/PDF）＋ プレゼン資料（PowerPoint/PDF・10MB未満）"],
        ["審査の流れ", "事前審査 → 一次審査（書類・11月末に通知）→ 最終審査（プレゼン）"],
        ["最終審査会", "2026-12-12（土）／和歌山県データ利活用推進センター"],
        ["プレゼン時間", "大学生部門 14分"],
        ["賞", "大賞20万円／データ利活用賞5万円／政策アイデア賞5万円（各部門1点）"],
    ], size=14, row_h=Emu(480000), head_h=Emu(360000), bold_col0=True)

    # ---------------------------------------------------------- 7 §3-2 審査観点
    s, y = slide(prs, "3-2 ／ 審査の観点", "この2点で審査されます（募集要項より）")
    half = int((CW - Emu(300000)) / 2)
    for i, (head, body) in enumerate([
        ("ア：データ利活用の観点",
         "独自の着眼点からデータを分析し、\n地域の現状や課題を抽出できているか"),
        ("イ：施策提案の観点",
         "独自性・将来性のある施策提案であり、\nデータ分析のエビデンスに基づいているか"),
    ]):
        cx = ML + i * (half + Emu(300000))
        card = _rect(s, cx, y, half, Emu(1600000), PALE, line_color=LINE)
        top = _rect(s, cx, y, half, Emu(80000), TEAL)
        top.line.fill.background()
        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Emu(280000)
        _para(tf, head, 19, bold=True, color=NAVY, space_after=14, first=True)
        _para(tf, body, 15, color=INK, space_after=0)

    y += Emu(1820000)
    y = note_box(
        s, y,
        "加えて「資料の完成度も審査します」と要項に明記されています。",
        h=Emu(560000), size=16,
    )

    y += Emu(200000)
    tf = _tb(s, ML, y, CW, Emu(300000))
    _para(tf, "（出典：第10回和歌山県データ利活用コンペティション 募集要項）",
          12, color=MUTED, space_after=0, first=True)

    # ---------------------------------------------------------- 8 §4 目的
    s, y = slide(prs, "4 ／ 分析の目的", None)
    box = _rect(s, ML, y + Emu(600000), CW, Emu(1700000), NAVY)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(400000)
    _para(tf, "和歌山県の住民が実際に感じている幸福度\n（主観的ウェルビーイング）を上げること。",
          30, bold=True, color=WHITE, space_after=0, first=True,
          align=PP_ALIGN.CENTER, line=1.4)

    # ---------------------------------------------------------- 9 §5 方法
    s, y = slide(prs, "5 ／ 目的を叶えるための方法", None)
    box = _rect(s, ML, y, CW, Emu(950000), TEAL)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _para(tf, "分析から政策提言までを行う。", 30, bold=True, color=WHITE,
          space_after=0, first=True, align=PP_ALIGN.CENTER)

    y += Emu(1200000)
    tf = _tb(s, ML, y, CW, Emu(400000))
    _para(tf, "理由", 15, bold=True, color=TEAL, space_after=0, first=True)
    bullets(s, y + Emu(460000), [
        "施策が明確になっていなければ、手は打てない",
        "手を打たなければ、幸福度は上がらない",
        "だから、分析だけでも提言だけでもなく、一続きで行う",
    ], size=19)

    # ---------------------------------------------------------- 10 §5 到達状態
    s, y = slide(prs, "5 ／ この方法をとった後の状態",
                 "次の4点を言える状態を目指します")
    y = table(s, ML, y, CW, [0.5, 5.4, 4.1], [
        ["#", "最終的に言えるようにすること", "それが言えると何ができるか"],
        ["1", "どの分野に手を打つべきか（優先順位と根拠）",
         "予算を振り向ける先が決まる"],
        ["2", "その分野で何を改善すべきか",
         "改善すべき対象が決まる"],
        ["3", "どこまでを目標とし、幸福度がどれだけ変わるか",
         "目標値と効果を定量的に示せる"],
        ["4", "そのために、どのような行政施策を実施すべきか",
         "費用対効果を踏まえた政策判断につなげられる"],
    ], size=15, row_h=Emu(620000))

    y += Emu(260000)
    note_box(
        s, y,
        "「どの分野の、何を、どこまで改善するために、どのような施策を実施すべきか」を言える。",
        h=Emu(620000), size=16,
    )

    # ---------------------------------------------------------- 11 §6-1/2 データ
    s, y = slide(prs, "6-1 ／ 使うデータ",
                 "デジタル庁「地域幸福度（Well-Being）指標」ダッシュボード")
    box = _rect(s, ML, y, CW, Emu(760000), PALE, line_color=LINE)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(280000)
    _para(tf, "市区町村単位で、住民の実感（主観）と地域の環境水準（客観）が"
              "同じ24分野で並ぶ。全国約530自治体分。",
          17, bold=True, color=NAVY, space_after=0, first=True)

    y += Emu(1000000)
    table(s, ML, y, CW, [3.4, 6.6], [
        ["取れるもの", "分かること"],
        ["幸福度（0〜10の総合スコア）", "住民がどれだけ幸せを感じているか　→ これを目的変数にする"],
        ["主観24分野（住民の実感）", "24分野それぞれを、住民が実際にどう感じているか"],
        ["客観24分野（環境の水準）", "同じ24分野で、地域の環境がどの水準にあるか"],
        ["客観の個別KPI（208本）", "客観24分野が、どの指標から成り立っているか"],
        ["年代別の内訳", "主観は全分野が年代別　→ 年代で実感がどう違うか"],
    ], size=15, row_h=Emu(490000), bold_col0=True)

    tf = _tb(s, ML, Emu(5700000), CW, Emu(500000))
    _para(tf, "「環境はあるのに、住民の実感が伴っていない分野」を特定できる。"
              "これが本分析の起点です。",
          17, bold=True, color=ORANGE, space_after=0, first=True)

    # ---------------------------------------------------------- 12 24分野
    s, y = slide(prs, "6-2 ／ 24分野の内訳", None)
    groups = [
        ("生活環境（16）", "医療・福祉／買物・飲食／住宅環境／移動・交通／遊び・娯楽／子育て／"
                        "初等・中等教育／地域行政／デジタル生活／公共空間／都市景観／自然景観／"
                        "自然の恵み／環境共生／自然災害／事故・犯罪"),
        ("地域の人間関係（2）", "地域とのつながり／多様性と寛容性"),
        ("自分らしい生き方（6）", "自己効力感／健康状態／文化・芸術／教育機会の豊かさ／雇用・所得／事業創造"),
    ]
    heights = [Emu(1300000), Emu(800000), Emu(950000)]
    cy = y + Emu(150000)
    for (head, body), h in zip(groups, heights):
        bar = _rect(s, ML, cy, Emu(50000), h, TEAL)
        bar.line.fill.background()
        card = _rect(s, ML + Emu(50000), cy, CW - Emu(50000), h, PALE,
                     line_color=None)
        tf = card.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Emu(280000)
        _para(tf, head, 17, bold=True, color=NAVY, space_after=10, first=True)
        _para(tf, body, 15, color=INK, space_after=0)
        cy += h + Emu(180000)

    # ---------------------------------------------------------- 13 §6-3 出所
    s, y = slide(prs, "6-3 ／ データの出所",
                 "主指標はすべて住民アンケート、客観指標は官庁統計からの合成")
    table(s, ML, y, CW, [2.6, 7.4], [
        ["区分", "出所"],
        ["主指標（主観）", "すべて住民アンケート。デジタル庁が全国の自治体で実施"],
        ["客観指標",
         "官庁統計から合成（国交省・厚労省・総務省）。\n原出典まで公開されている"],
        ["環境変数（統制用）",
         "総務省（人口・高齢化率・財政力指数）。\n自治体間の条件差を吸収するために外から足す"],
    ], size=15, row_h=Emu(680000), bold_col0=True)

    # ---------------------------------------------------------- 14 §6-4 分析単位
    s, y = slide(prs, "6-4 ／ 目的変数と分析単位", None)
    table(s, ML, y, CW, [2.4, 7.6], [
        ["項目", "内容"],
        ["目的変数", "幸福度（0〜10の生値）"],
        ["説明変数",
         "主観24分野・客観24分野（偏差値）、客観の個別KPI（実数）、\n"
         "環境変数（人口・高齢化率・財政力指数）"],
        ["分析単位", "1行 ＝ 1自治体 × 1年度 × 1年代"],
        ["対象",
         "推定は全国の自治体、診断の対象は和歌山。\n"
         "全国データで物差しを作り、それを和歌山に当てる"],
    ], size=15, row_h=Emu(700000), bold_col0=True)

    # ---------------------------------------------------------- 15 §7-1 用語
    s, y = slide(prs, "7-1 ／ この資料で使う言葉",
                 "分析設計の説明に入る前に、以下の5語を定義します")
    table(s, ML, y, CW, [2.6, 7.4], [
        ["言葉", "指すもの"],
        ["幸福度", "目的変数。住民に直接たずねた0〜10の総合スコア"],
        ["主観スコア", "24分野それぞれの、住民の実感の値（偏差値）"],
        ["客観スコア", "24分野それぞれの、地域の環境水準の値（偏差値）"],
        ["客観的構成値", "客観スコアを組み立てている個別KPI（208本）"],
        ["追加要因",
         "客観的構成値に含まれないが、主観スコアに効いている可能性のある要因\n"
         "（情報の届き方、予約のしやすさ、運用など）"],
    ], size=15, row_h=Emu(580000), bold_col0=True)

    # ---------------------------------------------------------- 16 §7-2 流れ
    s, y = slide(prs, "7-2 ／ 全体の流れ", "分析は9ステップですが、かたまりで見ると4つです")
    blocks = [
        ("①", "どの分野を扱うかを決める", "STEP 1・2", NAVY),
        ("②", "その分野の実感が、環境の数字で説明できるかを判定する", "STEP 3・4", TEAL),
        ("③", "何をどこまで動かすと、実感がどれだけ動くかを出す", "STEP 5・6・7", NAVY),
        ("④", "行政が打てる手に翻訳する", "STEP 8・9", TEAL),
    ]
    bh = Emu(810000)
    gap = Emu(300000)
    cy = y + Emu(120000)
    for i, (num, label, steps, color) in enumerate(blocks):
        chip = _rect(s, ML, cy, Emu(830000), bh, color)
        ctf = chip.text_frame
        ctf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _para(ctf, num, 26, bold=True, color=WHITE, space_after=0, first=True,
              align=PP_ALIGN.CENTER)
        body = _rect(s, ML + Emu(830000), cy, CW - Emu(830000) - Emu(1900000),
                     bh, PALE, line_color=None)
        btf = body.text_frame
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        btf.margin_left = Emu(320000)
        _para(btf, label, 19, bold=True, color=INK, space_after=0, first=True)
        tag = _rect(s, ML + CW - Emu(1900000), cy, Emu(1900000), bh, WHITE,
                    line_color=LINE)
        ttf = tag.text_frame
        ttf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _para(ttf, steps, 16, bold=True, color=color, space_after=0, first=True,
              align=PP_ALIGN.CENTER)
        if i < 3:
            ar = _rect(s, ML + Emu(300000), cy + bh + Emu(20000),
                       Emu(230000), gap - Emu(40000), LINE,
                       shape=MSO_SHAPE.DOWN_ARROW)
            ar.line.fill.background()
        if i == 1:
            note = _tb(s, ML + Emu(700000), cy + bh + Emu(40000),
                       Emu(4600000), Emu(240000))
            _para(note, "ここで進む道が2つに分かれる", 12, bold=True, color=ORANGE,
                  space_after=0, first=True)
        cy += bh + gap

    # ---------------------------------------------------------- 17 9ステップ前半
    s, y = slide(prs, "7-2 ／ 9ステップ（① 分野を絞る・② 道を分ける）", None)
    table(s, ML, y, CW, [1.3, 0.7, 4.6, 3.4], [
        ["", "STEP", "やること", "ここで分かること"],
        ["① 分野を絞る", "1", "24分野の主観スコアと幸福度の関連を見る",
         "幸福度と関連が大きい分野はどれか"],
        ["", "2", "重点カテゴリーを選定する", "和歌山が優先的に扱う分野はどれか"],
        ["② 道を分ける", "3", "重点カテゴリーの主観スコアを、\n客観的構成値で説明できるか確認する",
         "環境の数字で説明がつく分野か、\nつかない分野か"],
        ["", "4A", "【説明できる】客観的構成値を改善候補にする",
         "改善候補（既存のKPI）"],
        ["", "4B", "【説明できない】追加要因を探索して加える",
         "改善候補（既存のKPI ＋ 追加要因）"],
    ], size=14, row_h=Emu(640000), bold_col0=True)

    # ---------------------------------------------------------- 18 9ステップ後半
    s, y = slide(prs, "7-2 ／ 9ステップ（③ 効き目を見積もる・④ 施策にする）", None)
    table(s, ML, y, CW, [1.3, 0.7, 4.6, 3.4], [
        ["", "STEP", "やること", "ここで分かること"],
        ["③ 効き目を\n　 見積もる", "5", "伸ばすと実感が動きそうな候補に絞り込む",
         "効きそうな候補はどれか"],
        ["", "6", "和歌山に、伸ばす余地があるかを確認する", "実際に動かせる候補はどれか"],
        ["", "7", "候補に目標値を置き、主観スコアの変化を試算する",
         "どこまで動かすと、実感がどれだけ動くか"],
        ["④ 施策にする", "8", "行政施策の介入可能性と時間軸を整理する",
         "短期・中期・長期のどれで打てる手か"],
        ["", "9", "政策提案とロードマップを作成する", "和歌山が、何を、いつやるべきか"],
    ], size=14, row_h=Emu(620000), bold_col0=True)

    tf = _tb(s, ML, Emu(5950000), CW, Emu(300000))
    _para(tf, "※ 改善候補 ＝ STEP 4A では「客観的構成値」、"
              "STEP 4B では「客観的構成値」＋探索した「追加要因」",
          13, color=MUTED, space_after=0, first=True)

    # ---------------------------------------------------------- 19 §7-3 分岐
    s, y = slide(prs, "7-3 ／ 設計の要は STEP 3 の分岐",
                 "説明できるかどうかで、次の一手が変わる",
                 sub="重点カテゴリの主観スコアを、客観的構成値でどこまで説明できるかを確認します")
    y = table(s, ML, y, CW, [2.4, 3.6, 4.0], [
        ["STEP 3 の判定", "次にやること", "分析の目的"],
        ["十分に説明できる", "STEP 4A：個別KPI まで掘り下げる",
         "具体的な改善候補を特定する"],
        ["十分に説明できない", "STEP 4B：追加要因を探索する",
         "既存の指標では捉えられていない要因を\n見つけ、改善候補に加える"],
    ], size=15, row_h=Emu(780000), bold_col0=True)

    y += Emu(400000)
    note_box(
        s, y,
        "STEP 4B の追加要因の探索に、他チームとの差別化になりうる\n"
        "定性データ・地理データの活用が入ります。",
        h=Emu(700000), size=16,
    )

    # ---------------------------------------------------------- 20 クロージング
    s = _blank(prs)
    _rect(s, Emu(0), Emu(0), SW, SH, NAVY)
    _rect(s, Emu(0), SH - Emu(180000), SW, Emu(180000), TEAL)
    tf = _tb(s, ML, Emu(2400000), CW, Emu(1800000))
    _para(tf, "気になった点は、説明の途中でも遠慮なくおっしゃってください。",
          26, bold=True, color=WHITE, space_after=18, first=True,
          align=PP_ALIGN.CENTER)
    _para(tf, "「この前提が違う」「実務ではこうする」の一言が一番ありがたいです。",
          19, color=RGBColor(0x9F, 0xE0, 0xE8), space_after=0,
          align=PP_ALIGN.CENTER)

    footer(prs)
    return prs


if __name__ == "__main__":
    prs = build()
    prs.save(OUT)
    print(f"wrote {OUT.name}  ({len(prs.slides._sldIdLst)} slides)")
