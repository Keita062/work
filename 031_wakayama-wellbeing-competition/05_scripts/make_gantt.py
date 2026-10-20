# -*- coding: utf-8 -*-
"""和歌山コンペ WBS + ガントチャート生成
出典: 01_docs/和歌山コンペ_分析要件定義書 (1)0810.md の「wbs」セクション（2026-08-07 MTG）

v2 変更点:
  - タスクオーナー（B列）をプルダウン化（鈴木 / 西村 / 鈴木・西村 / 未定）
  - 進捗の手入力（0〜100%）を廃止し、F列「ステータス」プルダウン → G列「進捗率」自動計算に変更
  - ステータスの昇格条件・タスクごとの完了条件を「使い方と進捗の定義」シートに明記
"""
import datetime as dt
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "02_memo" / "和歌山コンペ_WBS_ガントチャート.xlsx"

D = dt.date

# ---------------------------------------------------------------- 配色
C_MONTH   = "1F3864"
C_WEEK    = "2E5F8A"
C_SECTION = "D9D9D9"
C_BAR     = "4A86C8"
C_DONE    = "1F4E79"
C_WKND    = "F2F2F2"
C_SUN     = "FCE4E4"
C_TODAY   = "FF6B6B"
C_MILE    = "FFF2CC"

# ステータスごとの色（F列の塗り分け）
ST_COLORS = {
    "未着手":     ("F2F2F2", "7F7F7F"),
    "着手":       ("DEEBF7", "1F4E79"),
    "作業中":     ("BDD7EE", "1F4E79"),
    "レビュー待ち": ("FFE699", "833C00"),
    "完了":       ("C6E0B4", "375623"),
}

thin = Side(style="thin",   color="BFBFBF")
hair = Side(style="hair",   color="D9D9D9")
med  = Side(style="medium", color="7F7F7F")

GUIDE = "使い方と進捗の定義"

# ---------------------------------------------------------------- ステータス定義
# (名称, 進捗率, この状態に上げる条件)
STATUS = [
    ("未着手", 0.00, "まだ手をつけていない。開始日を過ぎてもここなら遅延。"),
    ("着手", 0.25, "着手を宣言し、必要な材料（データ・資料・連絡先・参照リンク）を洗い出して手元に揃えた。"),
    ("作業中", 0.50, "成果物の初版がファイルとして存在し、リポジトリにコミットされている。中身の完成度は問わない。"),
    ("レビュー待ち", 0.75, "相手（西村氏／教授陣／竹本様／パーソルのDS等）に共有し、返答・確認を待っている。自分の手は離れている。"),
    ("完了", 1.00, "相手の確認が取れ、指摘の反映まで終わった。次のタスクにそのまま渡せる状態。"),
]

OWNERS = [
    ("鈴木", "鈴木が単独で持つタスク。"),
    ("西村", "西村氏が単独で持つタスク。"),
    ("鈴木・西村", "2人で一緒にやるタスク。どちらかが欠けると進まないもの。"),
    ("未定", "WBS原文に担当の明記がなく、まだ割り振っていないタスク。"),
]

# ---------------------------------------------------------------- タスク定義
# (kind, タイトル, オーナー, 開始日, 期限)
ROWS = [
    ("section", "本分析", "", None, None),

    ("task", "1. お互いに残論点について詰める ＆ 教授陣に確認してもらう",
     "鈴木・西村", D(2026, 8, 7), D(2026, 8, 10)),
    ("sub", "1-1. 【追加タスク】ガントチャート作成", "鈴木", D(2026, 8, 7), D(2026, 8, 10)),
    ("sub", "1-2. 【追加タスク】竹本様に連絡", "未定", D(2026, 8, 7), D(2026, 8, 10)),

    ("task", "2. 共有（22:00〜）", "鈴木・西村", D(2026, 8, 10), D(2026, 8, 10)),

    ("task", "3. パーソルキャリアの企画職・DSの方に見てもらう",
     "未定", D(2026, 8, 19), D(2026, 9, 4)),

    ("task", "4. 議論", "未定", None, None),

    ("task", "5. 6. データ格納 ＆ 実行", "未定", D(2026, 8, 23), D(2026, 9, 6)),
    ("sub", "5-1. リポジトリ格納：和歌山県 Q1 主観WB（幸福度／満足度）",
     "未定", D(2026, 8, 23), D(2026, 9, 6)),
    ("sub", "5-2. リポジトリ格納：和歌山県 主観WB 各カテゴリ",
     "未定", D(2026, 8, 23), D(2026, 9, 6)),
    ("sub", "5-3. リポジトリ格納：他県の主観／客観WB（時と場合による）",
     "未定", D(2026, 8, 23), D(2026, 9, 6)),

    ("task", "7. パワーポイントの作成 ＆ 応募用紙作成",
     "未定", D(2026, 9, 7), D(2026, 9, 16)),

    ("task", "8. 練習 ＆ 各個人で他者に見てもらう",
     "未定", D(2026, 9, 17), D(2026, 9, 27)),

    ("task", "9. FB改善", "未定", D(2026, 9, 28), D(2026, 9, 30)),

    ("task", "10. 最終書類確認 ＆ 提出", "未定", D(2026, 10, 1), D(2026, 10, 1)),
    ("milestone", "◆ 応募締切（10月8日 木）", "鈴木・西村", D(2026, 10, 8), D(2026, 10, 8)),

    ("task", "11. 練習", "未定", None, None),

    ("milestone", "◆ 12. 本番：和歌山県データ利活用推進センター 10:00〜",
     "鈴木・西村", D(2026, 12, 12), D(2026, 12, 12)),

    ("section", "応募書類の作成", "", None, None),
    ("task", "応募用紙", "未定", None, None),
    ("task", "プレゼン資料", "未定", None, None),
]

# タスクごとの「完了」の定義（DoD）
DOD = [
    ("1. お互いに残論点について詰める ＆ 教授陣に確認してもらう",
     "残論点それぞれに結論が出ていて、教授陣からのコメントを受け取り、要件定義書に反映済み。"),
    ("1-1. ガントチャート作成",
     "本ファイルが完成し、西村氏に共有して合意が取れている。"),
    ("1-2. 竹本様に連絡",
     "連絡を送信し、返信を受け取っている（送っただけはレビュー待ち止まり）。"),
    ("2. 共有（22:00〜）",
     "MTGを実施し、決まったことを要件定義書か Context Vault に書き込んだ。"),
    ("3. パーソルキャリアの企画職・DSの方に見てもらう",
     "面談を実施し、もらった指摘を一覧の形で書き出した。"),
    ("4. 議論",
     "3 の指摘それぞれに「対応する／しない」と理由を決め、要件定義書に反映した。"),
    ("5. 6. データ格納 ＆ 実行",
     "下の 5-1〜5-3 が揃い、分析スクリプトが最後まで通って結果ファイルが出力されている。"),
    ("5-1. 和歌山県 Q1 主観WB（幸福度／満足度）",
     "該当データを取得してリポジトリにコミットし、出典URLと取得日を記録した。"),
    ("5-2. 和歌山県 主観WB 各カテゴリ",
     "同上。24指標のうちどれを採用したかも記録する。"),
    ("5-3. 他県の主観／客観WB",
     "同上。使わないと判断した場合は、その判断理由を記録すれば完了とする。"),
    ("7. パワーポイントの作成 ＆ 応募用紙作成",
     "応募用紙が全項目埋まり、プレゼン資料が最初から最後まで通しでつながっている。"),
    ("8. 練習 ＆ 各個人で他者に見てもらう",
     "通しで練習し、第三者からのフィードバックを一覧に書き出した。"),
    ("9. FB改善",
     "フィードバック一覧の各項目に「反映済み」か「見送り＋理由」が入っている。"),
    ("10. 最終書類確認 ＆ 提出",
     "提出を完了し、受領の確認が取れている。"),
    ("11. 練習",
     "本番と同じ時間配分で通し練習を実施した。"),
    ("12. 本番",
     "発表を実施した。"),
    ("応募用紙 / プレゼン資料",
     "7 と同じ基準。7 の内訳としてこの2行で進捗を見る。"),
]

# ---------------------------------------------------------------- カレンダー範囲
CAL_START = D(2026, 8, 3)
CAL_END   = D(2026, 12, 13)
DAYS = [CAL_START + dt.timedelta(days=i) for i in range((CAL_END - CAL_START).days + 1)]
DOW_JP = "月火水木金土日"

FIRST_CAL_COL = 8            # H列からカレンダー（G列に進捗率を追加したため右へ1つ）
HDR_MONTH, HDR_WEEK, HDR_DAY, HDR_DOW = 1, 2, 3, 4
FIRST_DATA_ROW = 5

wb = Workbook()
ws = wb.active
ws.title = "WBS_ガントチャート"

# ---------------------------------------------------------------- 左側ヘッダ
LEFT = [
    ("タスクのタイトル", 46),
    ("タスクオーナー", 13),
    ("開始日", 11),
    ("期限", 11),
    ("期間", 6),
    ("ステータス", 14),
    ("進捗率", 8),
]
for i, (label, width) in enumerate(LEFT, start=1):
    col = get_column_letter(i)
    ws.column_dimensions[col].width = width
    ws.merge_cells(start_row=HDR_MONTH, start_column=i, end_row=HDR_DOW, end_column=i)
    c = ws.cell(row=HDR_MONTH, column=i, value=label)
    c.font = Font(bold=True, size=11, color="1F3864")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.fill = PatternFill("solid", fgColor="FFFFFF")
    c.border = Border(top=thin, bottom=med, left=thin, right=thin)

# ---------------------------------------------------------------- カレンダーヘッダ
run_start = 0
for idx in range(len(DAYS) + 1):
    end_of_run = (idx == len(DAYS)) or (DAYS[idx].month != DAYS[run_start].month)
    if end_of_run:
        c1 = FIRST_CAL_COL + run_start
        c2 = FIRST_CAL_COL + idx - 1
        ws.merge_cells(start_row=HDR_MONTH, start_column=c1, end_row=HDR_MONTH, end_column=c2)
        c = ws.cell(row=HDR_MONTH, column=c1,
                    value=f"{DAYS[run_start].year}年{DAYS[run_start].month}月")
        c.font = Font(bold=True, size=11, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=C_MONTH)
        c.alignment = Alignment(horizontal="center", vertical="center")
        run_start = idx

for w in range(len(DAYS) // 7):
    c1 = FIRST_CAL_COL + w * 7
    c2 = c1 + 6
    ws.merge_cells(start_row=HDR_WEEK, start_column=c1, end_row=HDR_WEEK, end_column=c2)
    c = ws.cell(row=HDR_WEEK, column=c1, value=f"第{w+1}週")
    c.font = Font(bold=True, size=10, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=C_WEEK)
    c.alignment = Alignment(horizontal="center", vertical="center")

for i, day in enumerate(DAYS):
    col = FIRST_CAL_COL + i
    ws.column_dimensions[get_column_letter(col)].width = 2.7

    cd = ws.cell(row=HDR_DAY, column=col, value=day)
    cd.number_format = "d"
    cd.font = Font(size=8, bold=True, color="1F3864")
    cd.alignment = Alignment(horizontal="center", vertical="center")
    cd.border = Border(left=hair, right=hair, top=thin)

    cw = ws.cell(row=HDR_DOW, column=col, value=DOW_JP[day.weekday()])
    cw.font = Font(size=8,
                   color="C00000" if day.weekday() == 6 else
                         ("0070C0" if day.weekday() == 5 else "404040"))
    cw.alignment = Alignment(horizontal="center", vertical="center")
    cw.border = Border(left=hair, right=hair, bottom=med)

    fill = C_SUN if day.weekday() == 6 else (C_WKND if day.weekday() == 5 else None)
    if fill:
        cd.fill = PatternFill("solid", fgColor=fill)
        cw.fill = PatternFill("solid", fgColor=fill)

ws.row_dimensions[HDR_MONTH].height = 18
ws.row_dimensions[HDR_WEEK].height = 16
ws.row_dimensions[HDR_DAY].height = 14
ws.row_dimensions[HDR_DOW].height = 14

# ---------------------------------------------------------------- データ行
LAST_CAL_COL = FIRST_CAL_COL + len(DAYS) - 1
r = FIRST_DATA_ROW

for kind, title, owner, start, end in ROWS:
    is_section = kind == "section"
    is_mile = kind == "milestone"

    indent = 0 if is_section else (1 if kind in ("task", "milestone") else 2)
    a = ws.cell(row=r, column=1, value=title)
    a.alignment = Alignment(horizontal="left", vertical="center", indent=indent)
    a.font = Font(bold=is_section or is_mile, size=10,
                  color="1F3864" if is_section else ("833C00" if is_mile else "000000"))

    ws.cell(row=r, column=2, value=owner if owner else None).alignment = \
        Alignment(horizontal="center", vertical="center")

    cs = ws.cell(row=r, column=3, value=start)
    ce = ws.cell(row=r, column=4, value=end)
    for c in (cs, ce):
        c.number_format = "yyyy/m/d"
        c.alignment = Alignment(horizontal="center", vertical="center")

    if not is_section:
        ws.cell(row=r, column=5,
                value=f'=IF(OR($C{r}="",$D{r}=""),"",$D{r}-$C{r}+1)')
        # F列: ステータス（プルダウン。初期値は未着手）
        ws.cell(row=r, column=6, value="未着手")
        # G列: 進捗率（F列から自動計算。手入力しない）
        cg = ws.cell(row=r, column=7,
                     value=f"=IFERROR(VLOOKUP($F{r},'{GUIDE}'!$A$12:$B$16,2,FALSE),0)")
        cg.number_format = "0%"
    for col in (5, 6, 7):
        ws.cell(row=r, column=col).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=6).font = Font(size=9)

    for col in range(1, LAST_CAL_COL + 1):
        cell = ws.cell(row=r, column=col)
        if is_section:
            cell.fill = PatternFill("solid", fgColor=C_SECTION)
        elif is_mile:
            cell.fill = PatternFill("solid", fgColor=C_MILE)
        if col < FIRST_CAL_COL:
            cell.border = Border(left=thin, right=thin, top=hair, bottom=hair)
        else:
            cell.border = Border(left=hair, right=hair, bottom=hair)

    ws.row_dimensions[r].height = 19
    r += 1

LAST_DATA_ROW = r - 1
CAL_RANGE = (f"{get_column_letter(FIRST_CAL_COL)}{FIRST_DATA_ROW}:"
             f"{get_column_letter(LAST_CAL_COL)}{LAST_DATA_ROW}")
FCL = get_column_letter(FIRST_CAL_COL)

# ---------------------------------------------------------------- 条件付き書式（ガント）
ws.conditional_formatting.add(
    CAL_RANGE,
    FormulaRule(formula=[f'{FCL}$3=TODAY()'],
                border=Border(left=Side(style="medium", color=C_TODAY),
                              right=Side(style="medium", color=C_TODAY)),
                stopIfTrue=False))

# 実績バー = 開始日から 期間×進捗率 日分（G列＝ステータス由来）
ws.conditional_formatting.add(
    CAL_RANGE,
    FormulaRule(
        formula=[f'AND($C5<>"",$D5<>"",$G5>0,{FCL}$3>=$C5,'
                 f'{FCL}$3<=$C5+ROUNDUP(($D5-$C5+1)*$G5,0)-1)'],
        fill=PatternFill(start_color=C_DONE, end_color=C_DONE, fill_type="solid"),
        stopIfTrue=True))

ws.conditional_formatting.add(
    CAL_RANGE,
    FormulaRule(
        formula=[f'AND($C5<>"",$D5<>"",{FCL}$3>=$C5,{FCL}$3<=$D5)'],
        fill=PatternFill(start_color=C_BAR, end_color=C_BAR, fill_type="solid"),
        stopIfTrue=True))

ws.conditional_formatting.add(
    CAL_RANGE,
    FormulaRule(formula=[f'WEEKDAY({FCL}$3,2)=7'],
                fill=PatternFill(start_color=C_SUN, end_color=C_SUN, fill_type="solid"),
                stopIfTrue=True))
ws.conditional_formatting.add(
    CAL_RANGE,
    FormulaRule(formula=[f'WEEKDAY({FCL}$3,2)=6'],
                fill=PatternFill(start_color=C_WKND, end_color=C_WKND, fill_type="solid"),
                stopIfTrue=True))

# ---------------------------------------------------------------- ステータス列の色分け
ST_RANGE = f"F{FIRST_DATA_ROW}:F{LAST_DATA_ROW}"
for name, (bg, fg) in ST_COLORS.items():
    ws.conditional_formatting.add(
        ST_RANGE,
        FormulaRule(formula=[f'$F{FIRST_DATA_ROW}="{name}"'],
                    fill=PatternFill(start_color=bg, end_color=bg, fill_type="solid"),
                    font=Font(color=fg, bold=(name == "完了"), size=9),
                    stopIfTrue=True))

# 進捗率のデータバー
ws.conditional_formatting.add(
    f"G{FIRST_DATA_ROW}:G{LAST_DATA_ROW}",
    DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                color="63BE7B", showValue=True, minLength=None, maxLength=None))

# ---------------------------------------------------------------- 入力規則（プルダウン）
dv_owner = DataValidation(
    type="list", formula1=f"='{GUIDE}'!$A$5:$A$8", allow_blank=True,
    showErrorMessage=True, errorTitle="タスクオーナー",
    error="鈴木 / 西村 / 鈴木・西村 / 未定 から選んでください。")
dv_owner.promptTitle = "タスクオーナー"
dv_owner.prompt = "担当を選択。2人でやるタスクは「鈴木・西村」。"
dv_owner.showInputMessage = True
ws.add_data_validation(dv_owner)

dv_status = DataValidation(
    type="list", formula1=f"='{GUIDE}'!$A$12:$A$16", allow_blank=True,
    showErrorMessage=True, errorTitle="ステータス",
    error="未着手 / 着手 / 作業中 / レビュー待ち / 完了 から選んでください。")
dv_status.promptTitle = "ステータス"
dv_status.prompt = ("選ぶと進捗率(G列)とガントの実績バーが自動で変わります。\n"
                    "各段階の条件は「使い方と進捗の定義」シートを参照。")
dv_status.showInputMessage = True
ws.add_data_validation(dv_status)

rr = FIRST_DATA_ROW
for kind, title, owner, start, end in ROWS:
    if kind != "section":
        dv_owner.add(f"B{rr}")
        dv_status.add(f"F{rr}")
    rr += 1

# ---------------------------------------------------------------- 仕上げ
ws.freeze_panes = f"{FCL}{FIRST_DATA_ROW}"
ws.sheet_view.showGridLines = False
ws.auto_filter.ref = f"A{HDR_DOW}:G{LAST_DATA_ROW}"

ws.print_title_rows = f"{HDR_MONTH}:{HDR_DOW}"
ws.print_title_cols = "$A:$G"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# 凡例
lg = LAST_DATA_ROW + 2
ws.cell(row=lg, column=1, value="凡例").font = Font(bold=True, size=10, color="1F3864")
legend = [
    (C_BAR,     "予定期間"),
    (C_DONE,    "実績（ステータスに連動して自動で伸びる）"),
    (C_MILE,    "マイルストーン（締切・本番）"),
    (C_SECTION, "大分類"),
    (C_WKND,    "土曜／日曜"),
]
for i, (color, text) in enumerate(legend):
    row_ = lg + 1 + i
    sw = ws.cell(row=row_, column=1, value="")
    sw.fill = PatternFill("solid", fgColor=color)
    sw.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.cell(row=row_, column=2, value=text).alignment = Alignment(horizontal="left")

note_row = lg + len(legend) + 2
notes = [
    "出典：01_docs/和歌山コンペ_分析要件定義書 (1)0810.md「wbs」セクション（2026年8月7日 MTG）",
    "B列（タスクオーナー）とF列（ステータス）はプルダウン。セルを選ぶと右に▼が出ます。",
    "G列（進捗率）は F列から自動計算される数式です。直接入力しないでください。",
    "各ステータスに上げる条件と、タスクごとの「完了」の定義は「使い方と進捗の定義」シートにあります。",
    "「4. 議論」「11. 練習」「応募用紙」「プレゼン資料」は原文に日付が無いため空欄。日付を入れるとバーが自動で描かれます。",
]
for i, t in enumerate(notes):
    c = ws.cell(row=note_row + i, column=1, value=("※ " + t))
    c.font = Font(size=9, color="595959")

# ================================================================ 定義シート
gs = wb.create_sheet(GUIDE)
gs.sheet_view.showGridLines = False
for col, wdt in zip("ABC", (22, 10, 78)):
    gs.column_dimensions[col].width = wdt

def head(row, text, span=3):
    gs.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    c = gs.cell(row=row, column=1, value=text)
    c.font = Font(bold=True, size=12, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=C_MONTH)
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    gs.row_dimensions[row].height = 22

def thead(row, labels):
    for i, t in enumerate(labels, start=1):
        c = gs.cell(row=row, column=i, value=t)
        c.font = Font(bold=True, size=10, color="1F3864")
        c.fill = PatternFill("solid", fgColor="DCE6F1")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(top=thin, bottom=med, left=thin, right=thin)
    gs.row_dimensions[row].height = 18

t = gs.cell(row=1, column=1, value="使い方と進捗の定義")
t.font = Font(bold=True, size=16, color="1F3864")
gs.row_dimensions[1].height = 26
gs.cell(row=2, column=1,
        value="このシートの表が、そのままガントチャート側のプルダウンの選択肢になっています。"
              "文言を書き換えると選択肢も変わります。").font = Font(size=9, color="595959")

# --- 1. タスクオーナー（A5:A8 が入力規則の参照元）
head(3, "1. タスクオーナー（ガントチャートのB列）")
thead(4, ["選択肢", "", "意味"])
gs.merge_cells(start_row=4, start_column=2, end_row=4, end_column=2)
for i, (name, desc) in enumerate(OWNERS):
    row_ = 5 + i
    c1 = gs.cell(row=row_, column=1, value=name)
    c1.font = Font(bold=True, size=10)
    c1.alignment = Alignment(horizontal="center", vertical="center")
    c3 = gs.cell(row=row_, column=3, value=desc)
    c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c3.font = Font(size=9)
    for col in range(1, 4):
        gs.cell(row=row_, column=col).border = Border(left=thin, right=thin,
                                                      top=hair, bottom=hair)
    gs.row_dimensions[row_].height = 17

# --- 2. ステータス（A12:A16 が選択肢、A12:B16 が進捗率のVLOOKUP参照元）
head(10, "2. ステータス（F列）と進捗率（G列）")
thead(11, ["ステータス", "進捗率", "この状態に上げてよい条件（＝何をしたら次に進むか）"])
for i, (name, ratio, cond) in enumerate(STATUS):
    row_ = 12 + i
    bg, fg = ST_COLORS[name]
    c1 = gs.cell(row=row_, column=1, value=name)
    c1.font = Font(bold=True, size=10, color=fg)
    c1.fill = PatternFill("solid", fgColor=bg)
    c1.alignment = Alignment(horizontal="center", vertical="center")

    c2 = gs.cell(row=row_, column=2, value=ratio)
    c2.number_format = "0%"
    c2.font = Font(bold=True, size=10)
    c2.alignment = Alignment(horizontal="center", vertical="center")

    c3 = gs.cell(row=row_, column=3, value=cond)
    c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c3.font = Font(size=9)
    for col in range(1, 4):
        gs.cell(row=row_, column=col).border = Border(left=thin, right=thin,
                                                      top=hair, bottom=hair)
    gs.row_dimensions[row_].height = 30

gs.cell(row=17, column=1,
        value="※ 進捗率(G列)はこの表の「進捗率」を引いてくる数式です。数字を直接いじらず、ステータスを選び直してください。"
        ).font = Font(size=9, color="C00000")
gs.cell(row=18, column=1,
        value="※ 迷ったら下げる。「レビュー待ち」は相手ボールのときだけ。自分の手元に戻ってきたら「作業中」に戻します。"
        ).font = Font(size=9, color="595959")

# --- 3. タスクごとの完了の定義
head(20, "3. タスクごとの「完了（100%）」の定義")
thead(21, ["タスク", "", "これが満たせたら「完了」"])
gs.merge_cells(start_row=21, start_column=1, end_row=21, end_column=2)
for i, (task, cond) in enumerate(DOD):
    row_ = 22 + i
    gs.merge_cells(start_row=row_, start_column=1, end_row=row_, end_column=2)
    c1 = gs.cell(row=row_, column=1, value=task)
    c1.font = Font(size=9, bold=True)
    c1.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c3 = gs.cell(row=row_, column=3, value=cond)
    c3.font = Font(size=9)
    c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for col in range(1, 4):
        gs.cell(row=row_, column=col).border = Border(left=thin, right=thin,
                                                      top=hair, bottom=hair)
    gs.row_dimensions[row_].height = 28

wb.save(OUT)
print("saved:", OUT)
print("data rows:", FIRST_DATA_ROW, "-", LAST_DATA_ROW, "| cal cols:", len(DAYS),
      "| guide rows:", 22 + len(DOD) - 1)
