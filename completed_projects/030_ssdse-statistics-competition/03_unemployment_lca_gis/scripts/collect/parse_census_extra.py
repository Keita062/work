#!/usr/bin/env python3
"""
令和2年（2020）国勢調査 就業状態等基本集計の追加整形。

入力（DL不要・既にローカル）:
  data/raw/estat/census_c06_03_age_industry.xlsx
      第6-3表 男女・年齢(5歳階級)・産業(大分類)別 就業者数。
      列: 0=地域識別コード 1=都道府県 2=地域名 3=男女 4=産業3部門 5=産業
          6=就業者数(00_総数) 7..23=年齢5歳階級 24=R1(15-64) 25=R2(65歳以上) ...
      地域識別コード 'a' = 全国＋47都道府県（市区町村は 0/1/2/3）。
  data/raw/estat/census_c04_laborforce_industry.xlsx
      第4表。列7-11 は全て「就業者数」を労働力状態(総数/主に仕事/家事のほか仕事/
      通学のかたわら仕事/休業者)で細分したもの。完全失業者・労働力人口の列は無い。
      → 県×産業の完全失業者/労働力人口/完全失業率は本表からは算出不能（理由は下記）。

出力:
  data/processed/census_age_industry_pref.csv
      prefecture, industry(9業種), 高齢就業者比率, 若年就業者比率, year=2020

実行:
  PYTHONUTF8=1 PYTHONIOENCODING=utf-8 python scripts/collect/parse_census_extra.py
"""
import csv
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent.parent
RAW = ROOT / "data" / "raw" / "estat"
PROC = ROOT / "data" / "processed"
PROC.mkdir(parents=True, exist_ok=True)

C06 = RAW / "census_c06_03_age_industry.xlsx"
C04 = RAW / "census_c04_laborforce_industry.xlsx"

# 産業大分類(1文字) → 9業種（プロジェクト既存の命名に一致：全角読点「，」/中黒「・」）
LETTER_TO_GROUP = {
    "E": "製造業",
    "I": "卸売業，小売業",
    "F": "サービス業他", "L": "サービス業他", "M": "サービス業他",
    "N": "サービス業他", "O": "サービス業他", "P": "サービス業他",
    "Q": "サービス業他", "R": "サービス業他",
    "A": "農・林・漁・鉱業", "B": "農・林・漁・鉱業", "C": "農・林・漁・鉱業",
    "D": "建設業",
    "G": "情報通信業",
    "H": "運輸業",
    "J": "金融・保険業",
    "K": "不動産業",
}
GROUP_ORDER = ["製造業", "卸売業，小売業", "サービス業他", "農・林・漁・鉱業",
               "建設業", "情報通信業", "運輸業", "金融・保険業", "不動産業"]

# c06_03 の年齢列インデックス（0起点）
COL_TOTAL = 6            # 00_総数（就業者総数）
COL_YOUNG = [7, 8, 9, 10]  # 01_15~19, 02_20~24, 03_25~29, 04_30~34 → 15~34歳
COL_ELDER = 25          # R2_（再掲）65歳以上

MAJOR_RE = re.compile(r"^([A-R])_")  # A_～R_ の大分類のみ（"01_うち..." や "R1_再掲" は除外）


def to_num(v):
    """国勢調査セル値を数値化。'-'（皆無）や空欄は0、'X'（秘匿）はNoneで返す。"""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "…", "***"):
        return 0.0
    if s in ("X", "x", "－"):  # 秘匿など
        return None
    try:
        return float(s)
    except ValueError:
        return None


def pref_fullname(code_label: str) -> str:
    """'02_青森県' → '青森県'。"""
    return code_label.split("_", 1)[1] if "_" in code_label else code_label


def parse_c06():
    wb = openpyxl.load_workbook(C06, read_only=True, data_only=True)
    ws = wb["c06_03"]
    # agg[pref][group] = [total, young, elder]
    agg = {}
    suppressed = 0
    for row in ws.iter_rows(min_row=11, values_only=True):
        if row[0] != "a":            # 全国＋都道府県のみ
            continue
        if row[3] != "0_総数":         # 男女=総数
            continue
        pref_label = row[1]
        if pref_label == "00_全国":    # 全国は除外（都道府県のみ出力）
            continue
        m = MAJOR_RE.match(str(row[5]))
        if not m:                     # 大分類 A_～R_ 以外（総数/サブ項目/再掲/S/T）を除外
            continue
        letter = m.group(1)
        group = LETTER_TO_GROUP.get(letter)
        if group is None:             # S(公務)/T(分類不能) は9業種に含めない
            continue
        total = to_num(row[COL_TOTAL])
        elder = to_num(row[COL_ELDER])
        young_vals = [to_num(row[c]) for c in COL_YOUNG]
        if total is None or elder is None or any(v is None for v in young_vals):
            suppressed += 1
            continue
        young = sum(young_vals)
        pref = pref_fullname(pref_label)
        g = agg.setdefault(pref, {})
        acc = g.setdefault(group, [0.0, 0.0, 0.0])
        acc[0] += total
        acc[1] += young
        acc[2] += elder
    wb.close()

    rows = []
    for pref in agg:
        for group in GROUP_ORDER:
            if group not in agg[pref]:
                continue
            total, young, elder = agg[pref][group]
            elder_ratio = round(elder / total * 100, 4) if total > 0 else ""
            young_ratio = round(young / total * 100, 4) if total > 0 else ""
            rows.append((pref, group, elder_ratio, young_ratio, 2020))
    if suppressed:
        print(f"[c06] 秘匿等でスキップしたセル: {suppressed}", file=sys.stderr)
    return rows


def inspect_c04():
    """c04 が完全失業者/労働力人口を含むか機械的に確認（ヘッダの値列区分を判定）。"""
    wb = openpyxl.load_workbook(C04, read_only=True, data_only=True)
    ws = wb["c04"]
    header = None
    for i, row in enumerate(ws.iter_rows(min_row=8, max_row=8, values_only=True)):
        header = row
    wb.close()
    labels = [str(v) for v in header[6:] if v is not None]
    has_unemp = any(("失業" in x) or ("労働力人口" in x) for x in labels)
    return labels, has_unemp


def main():
    rows = parse_c06()
    out = PROC / "census_age_industry_pref.csv"
    with out.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["prefecture", "industry", "高齢就業者比率", "若年就業者比率", "year"])
        w.writerows(rows)
    prefs = sorted({r[0] for r in rows})
    inds = sorted({r[1] for r in rows})
    print(f"[出力] {out}")
    print(f"  行数={len(rows)} 都道府県数={len(prefs)} 産業数={len(inds)}")

    labels, has_unemp = inspect_c04()
    print(f"[c04] 値列区分={labels}")
    print(f"[c04] 完全失業者/労働力人口の列: {'あり' if has_unemp else 'なし'}"
          " → 県×産業の完全失業率は本表からは算出不能")


if __name__ == "__main__":
    main()
